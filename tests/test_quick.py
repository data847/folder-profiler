"""Step 1 - the quick count: same numbers as the full profile, but no file is opened.   Run: python tests/test_quick.py"""
import builtins, csv, io, os, sys, tempfile, zipfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT); sys.path.insert(0, os.path.join(ROOT, "claude-skill", "profile-compile", "scripts"))
import file_profiler as fp
bad = []
def check(label, got, want):
    ok = got == want; print(("ok   " if ok else "FAIL ") + f"{label}: {got}" + ("" if ok else f"   (wanted {want})")); bad.append(not ok)
rd = lambda z, n: list(csv.DictReader(io.StringIO(zipfile.ZipFile(z).read(n).decode("utf-8-sig"))))
def put(base, rel, body=b"x"):
    path = os.path.join(base, rel); os.makedirs(os.path.dirname(path), exist_ok=True); open(path, "wb").write(body)

with tempfile.TemporaryDirectory() as t:
    src = os.path.join(t, "data")
    for i in range(30):
        put(src, f"Alpha/sub{i % 3}/f{i}.{['sldprt', 'sldasm', 'pdf', 'DWG'][i % 4]}", b"x" * (i + 1))
    for i in range(10):
        put(src, f"Beta/deep/er/still/f{i}.step", b"y" * (i + 5))
    put(src, "Beta/.DS_Store", b"j"); put(src, "Beta/Thumbs.db", b"j"); put(src, "Beta/._hidden", b"j")
    put(src, "Beta/bracket.prt.7", b"creo")                                  # Creo version number is not an extension
    put(src, "loose.txt", b"top-level file")                                  # a file directly in the chosen folder
    with zipfile.ZipFile(os.path.join(src, "Beta", "bundle.zip"), "w") as z:
        z.writestr("inside.sldprt", b"a" * 20); z.writestr("inside2.sldprt", b"b" * 20)
    deep = os.path.join(src, "Gamma")
    for i in range(150):
        deep = os.path.join(deep, f"l{i}")
    put(deep, "bottom.sldprt", b"z")
    if hasattr(os, "symlink"):
        try:
            os.symlink(src, os.path.join(src, "Alpha", "loop"), target_is_directory=True)     # a folder link that points back up
        except OSError:
            pass

    # ---- no file is opened
    real_open = builtins.open
    opened = []
    def guard(path, *a, **k):
        if str(path).startswith(src):
            opened.append(str(path))
        return real_open(path, *a, **k)
    builtins.open = guard
    try:
        qz, nproj, nfiles, errors = fp.quick_scan(src, os.path.join(t, "q"), log=lambda *_: None)
    finally:
        builtins.open = real_open
    check("no file of the data folder was opened", opened, [])
    info = {r["key"]: r["value"] for r in rd(qz, "Run_info.csv")}
    check("Run_info says quick count, nothing opened, depth beyond 150 reached", (info["mode"].startswith("quick"), info["files_opened"], int(info["deepest_folder_level"]) >= 150), (True, "0", True))

    # ---- same numbers as the full profile (level 0); neither step looks inside the zip
    fz, *_ = fp.profile(src, os.path.join(t, "f"), log=lambda *_: None, workers=2, system_info=False, tier=0)
    key = lambda z: {(r["project"], r["ext"]): (int(r["file_count"]), int(r["total_bytes"])) for r in rd(z, "Extensions.csv")}
    check("per-project extension counts and bytes equal the full profile's", key(qz) == key(fz), True)
    check("same projects and the same file total", (sorted(r["project"] for r in rd(qz, "Projects.csv")) == sorted(r["project"] for r in rd(fz, "Projects.csv")),
          int(info["files"]) == int(rd(fz, "Run_info.csv")[0 + [r["key"] for r in rd(fz, "Run_info.csv")].index("files")]["value"])), (True, True))
    allx = {r["ext"]: int(r["file_count"]) for r in rd(qz, "Extensions_All.csv")}
    check("Extensions_All adds the projects up; extensions are lower-case; Creo part.prt.7 counts as .prt", (allx["sldprt"] >= 8, "dwg" in allx, allx["prt"]), (True, True, 1))
    check("system files (.DS_Store, Thumbs.db, ._x) are not counted", ("ds_store" in allx or "db" in allx or "(none)" in allx), False)
    check("a zip is counted as one .zip file - nothing looked inside", (allx["zip"], "sldprt" in allx and allx["sldprt"] == sum(1 for i in range(30) if i % 4 == 0) + 1), (1, True))
    if hasattr(os, "symlink") and os.path.islink(os.path.join(src, "Alpha", "loop")):
        check("a folder link back up is not followed (counted, not looped)", int(info["linked_folders_skipped"]) >= 1, True)

    # ---- the command line: --quick is step 1, no flag is step 2
    import subprocess
    r = subprocess.run([sys.executable, os.path.join(ROOT, "file_profiler.py"), src, "--quick", "--out", os.path.join(t, "cli")], capture_output=True, text=True)
    check("command line --quick writes a quick_*.zip and says what to send", (r.returncode, [f for f in os.listdir(os.path.join(t, "cli")) if f.startswith("quick_")] != [], "ONE file" in r.stdout), (0, True, True))
    r = subprocess.run([sys.executable, os.path.join(ROOT, "file_profiler.py"), src, "--archives", "off"], capture_output=True, text=True)
    check("the old --archives option is gone (refused)", r.returncode != 0, True)
    # ---- picked folders mode
    qz2, n2, f2, _e = fp.quick_scan(src, os.path.join(t, "q2"), log=lambda *_: None, folders=[os.path.join(src, "Alpha"), os.path.join(src, "Beta")])
    check("picked folders: two projects, their own files only", (n2, sorted(r["project"] for r in rd(qz2, "Projects.csv"))), (2, ["Alpha", "Beta"]))
    # ---- project depth 2
    qz3, n3, f3, _e = fp.quick_scan(src, os.path.join(t, "q3"), log=lambda *_: None, project_depth=2)
    check("project depth 2 splits Alpha into its sub-folders", {r["project"] for r in rd(qz3, "Projects.csv")} >= {"Alpha/sub0", "Alpha/sub1", "Alpha/sub2"}, True)

    # ---- skill: step 1 zip -> workbook
    try:
        import compile_profiles as cp
    except ImportError:
        print("SKIP - compiler not present")
    else:
        em, sf = cp.load_mapping(os.path.join(ROOT, "claude-skill", "profile-compile", "scripts", "mapping"))
        out = cp.Path(t) / "qc"; out.mkdir()
        check("quick zip is recognised, a full profile zip is not", (cp.is_quick_zip(qz), cp.is_quick_zip(fz)), (True, False))
        np_, nf, cls_rows = cp.compile_quick([qz], em, sf, out)
        check("compiled overview adds up to the file count", (nf, sum(r[2] for r in cls_rows)), (int(info["files"]), int(info["files"])))
        check("files written", sorted(x.name for x in out.iterdir()), ["Quick_By_Extension.csv", "Quick_By_Project.csv", "Quick_Overview.csv", "quick_count.xlsx"] if cp.Workbook else ["Quick_By_Extension.csv", "Quick_By_Project.csv", "Quick_Overview.csv"])

# ---- unexpected extensions never break the count, and never become spreadsheet formulas
with tempfile.TemporaryDirectory() as t:
    odd = os.path.join(t, "odd", "P")
    names = ["noext", "trailingdot.", "Mixed.CaSe", ".gitignore", "evil.=1+1", "at.@sum", "dash.-5", "comma,name.c,d", 'quote".q"x',
             "tab\there.txt", "long." + "x" * 120, "unicode.\u6570\u636e", "UPPER.DWG", "dup.dwg", "spaces.my ext"]
    for n in names:
        put(odd, n, b"x" * 10)
    qz, _p, nfiles, _e = fp.quick_scan(os.path.join(t, "odd"), os.path.join(t, "q"), log=lambda *_: None)
    allx = {r["ext"]: int(r["file_count"]) for r in rd(qz, "Extensions_All.csv")}
    check("every odd file is counted (none lost, none crash)", (nfiles, sum(allx.values())), (len(names), len(names)))
    check("no extension -> (none); trailing dot -> (none); case folded; dotfile -> (none)", (allx["(none)"], allx["case"], allx["dwg"]), (3, 1, 2))
    check("commas, quotes, long and unicode extensions survive the CSV", ("c,d" in allx, 'q"x' in allx, "x" * 120 in allx, "\u6570\u636e" in allx), (True,) * 4)
    try:
        import compile_profiles as cp
    except ImportError:
        cp = None
    if cp is not None and cp.Workbook:
        import openpyxl
        em, sf = cp.load_mapping(os.path.join(ROOT, "claude-skill", "profile-compile", "scripts", "mapping"))
        out = cp.Path(t) / "oc"; out.mkdir()
        cp.compile_quick([qz], em, sf, out)
        ws = openpyxl.load_workbook(out / "quick_count.xlsx")["By Extension"]
        cells = {c.value: c for c in (r[0] for r in ws.iter_rows(min_row=2))}
        check("'=1+1' is stored as text, not as a formula", (cells["=1+1"].data_type, cells["=1+1"].value), ("s", "=1+1"))
        txt = open(out / "Quick_By_Extension.csv", encoding="utf-8-sig").read()
        check("in the CSV a formula-like extension gets a leading apostrophe", ("\n'=1+1," in txt and "\n'@sum," in txt and "\n-5," in txt), True)
        check("unknown extensions are listed as UNMAPPED with their counts (nothing dropped)", [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2) if r[0].value == "xxxx"[:0] or r[0].value == "my ext"], [("my ext", "UNMAPPED")])
sys.exit(1 if any(bad) else 0)
